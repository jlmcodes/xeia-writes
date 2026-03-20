import streamlit as st
import docx
import re
import uuid
import pandas as pd
from spellchecker import SpellChecker
from fpdf import FPDF

# --- Page Configuration & Strict Layout Control ---
st.set_page_config(page_title="Xeia Writes", page_icon="logo.png", layout="wide", initial_sidebar_state="collapsed")

# --- Callbacks & Helpers ---
def ignore_lapse(lapse_id, category_key): 
    st.session_state.ignored_lapses.add(lapse_id)
    st.session_state.open_lapses_category = category_key

def clean_for_pdf(text):
    replacements = { '“': '"', '”': '"', '‘': "'", '’': "'", '–': '-', '—': '-', '…': '...' }
    for search, replace in replacements.items():
        text = text.replace(search, replace)
    return text.encode('latin-1', 'ignore').decode('latin-1')

def get_smart_snippet(text, match_start=None, match_end=None, highlight_word=None):
    if not text: return ""
    if match_start is None or match_end is None:
        return text[:80] + "..." if len(text) > 80 else text
        
    if len(text) <= 90: 
        snippet = text
    else:
        start_chunk = text[:25]
        ctx_start = max(0, match_start - 25)
        ctx_end = min(len(text), match_end + 25)
        
        if ctx_start > 30: snippet = start_chunk.strip() + "... " + text[ctx_start:ctx_end].strip() + "..."
        else: snippet = text[:max(60, ctx_end)].strip() + "..."
        
    if highlight_word:
        pattern = re.compile(re.escape(highlight_word), re.IGNORECASE)
        snippet = pattern.sub(f"<span style='color:#AEA743; font-weight:700;'>{highlight_word}</span>", snippet)
        
    return snippet

# --- Initialization & Lexicons ---
spell = SpellChecker()

xeia_lexicon = {
    r'(?:^|\.\s+)(Also|And)\b': ("Moreover / Furthermore", "Adding a point that is more important than the previous one.", "If listing basic items, retain 'And/Also'. If adding a stronger supporting argument to your thesis, upgrade to 'Moreover' or 'Furthermore'."),
    r'(?:^|\.\s+)(So)\b': ("Therefore / Hence", "Showing a logical result.", "If used casually, upgrade to 'Therefore'. If functioning as 'so that', replace with 'in order to'."),
    r'(?:^|\.\s+)(But)\b': ("Conversely / However", "Introducing a direct opposite.", "Avoid starting academic sentences with 'But'. Upgrade to 'Conversely' or 'However' to maintain a formal tone."),
    r'\b(Unlike)\b': ("In contrast with", "Comparing two things to highlight differences.", "For a formal comparative analysis, 'In contrast with' carries a stronger academic weight than 'Unlike'."),
    r'\b(If not)\b': ("Otherwise", "Explaining conditional failures.", "If denoting a conditional failure (e.g., 'do this, if not...'), 'Otherwise' is far more precise."),
    r'\b(Even though)\b': ("Despite / Although", "Defying expectations.", "If preceding a noun phrase, use 'Despite'. If preceding a full clause with a subject and verb, 'Although' is better."),
    r'\b(Similarly)\b': ("In line with this", "Connecting to a previous theme.", "'In line with this' works best when connecting a new finding to a previously established academic rule."),
    r'\b(About)\b': ("As to / Regarding", "Introducing a specific topic.", "When introducing a topic formally, 'Regarding' or 'As to' is preferred over the casual 'About'."),
    r'\b(By this)\b': ("Hereby", "Formally declaring an action.", "'Hereby' is strictly for formal declarations or legal bindings. Use with caution."),
    r'\b(On it)\b': ("Thereon", "Referring to a specific place.", "'Thereon' is highly formal; ensure it refers to a specific physical or conceptual place just mentioned in the text."),
    r'\b(Related)\b': ("Corresponding", "Linking equivalent items.", "'Corresponding' implies a direct matching relationship, not just a casual link."),
    r'\b(Covers)\b': ("Encompasses", "Describing a wide range.", "'Encompasses' is ideal when describing the scope, limitations, or boundaries of a study."),
    r'\b(Using)\b': ("By means of / Utilizing", "Explaining a method or tool.", "If functioning as an active gerund (e.g., 'prohibited from using materials'), retain 'using'. If describing an analytical tool or methodology, upgrade to 'By means of' or 'Utilizing'."),
    r'\b(Clearly)\b': ("Evidently", "Highlighting a proven fact.", "'Evidently' is preferred in research when the claim is objectively backed by the data just presented."),
    r'\b(Take note)\b': ("It should be noted", "Drawing the reader's attention.", "'It should be noted' shifts the tone from a casual command to an objective academic observation."),
    r'\b(As per)\b': ("According to", "Attributing a specific claim.", "'According to' is the standard APA preference for citing authors, studies, or policies."),
    r'\b(Following)\b': ("In accordance with", "Confirming procedural alignment.", "'In accordance with' is best used for methodologies that strictly obey a set standard or law."),
    r'\b(Because of)\b': ("Based on the study of", "Rooting findings in evidence.", "When attributing a specific cause to literature, 'Based on the study of' is much more authoritative."),
    r'\b(The study shows)\b': ("It is found in the study of", "Pointing out a discovery.", "'It is found in the study of' or 'Research indicates' removes the anthropomorphic tone of a study 'showing' something."),
    r'\b(It says in)\b': ("It is mentioned in", "Referencing a supporting detail.", "Never use 'It says' for literature. Upgrade to 'It is mentioned in' or 'As noted by'.")
}

# --- Feasibility Study Master Syllabus ---
FEASIBILITY_OUTLINE = {
    "CHAPTER 1": "LEGAL AND TAXATION ASPECT", "1.1": "Legal Aspect", "1.1.1": "Securities and Exchange Commission",
    "1.1.2": "Mayor's Permit", "1.1.2.A": "Barangay Clearance", "1.1.2.B": "Locational Clearance",
    "1.1.2.C": "Occupancy Permit", "1.1.2.D": "Sanitary Permit", "1.1.2.E": "Fire, Safety and Inspection Certificate",
    "1.1.3": "Registration of Mandatory Government Agencies", "1.1.3.A": "Social Security System",
    "1.1.3.B": "Philhealth Insurance Corporation", "1.1.3.C": "Pag-Ibig Fund", "1.1.3.D": "Bureau of Internal Revenue",
    "1.2": "Taxation Aspect", "1.2.1": "Corporate Tax", "1.2.2": "Minimum Corporate Income Tax",
    "1.2.3": "Documentary Stamp Tax", "1.2.4": "Community Tax Certificate", "1.2.5": "Value Added Tax",
    "CHAPTER 2": "MANAGEMENT ASPECT", "2.1": "Business Profile", "2.1.1": "Business Name", "2.1.2": "Company Tagline",
    "2.1.3": "Company Logo", "2.2": "Type of Business Organization", "2.3": "The Incorporators", "2.4": "Management Style",
    "2.5": "Mission Statement", "2.6": "Vision Statement", "2.7": "Our Values", "2.8": "Objectives of the Company",
    "2.9": "Organizational Chart", "2.10": "Job Description and Qualifications", "2.10.1": "General Manager",
    "2.10.2": "Secretary", "2.10.3": "Marketing Manager", "2.10.4": "Marketing Staff", "2.10.5": "Operations Manager",
    "2.10.6": "Operations Staff", "2.10.7": "Finance Manager", "2.10.9": "Accounting Staff", "2.10.10": "Cashier",
    "2.11": "Projected Salary", "2.12": "Business Policy", "2.12.1": "Hiring Process", "2.12.2": "Training Policy",
    "2.12.3": "Performance and Evaluation Policy", "2.12.4": "Promotion Policy", "2.12.5": "Work Schedule & Dress Code Policy",
    "2.12.6": "Salaries and Wages/Thirteenth Month Pay Policy", "2.12.7": "Contract and Confidentiality Agreements",
    "2.12.8": "Leave of Absence", "2.12.9": "Attendance Policy", "2.12.10": "Recognition Policy",
    "2.12.11": "Termination or Resignation of Employee", "2.12.12": "Employee Benefits", 
    "2.12.13": "Illegal Drugs and Alcohol/Smoking/Telephone & Computer Use Policy", "2.12.14": "Employment Classification",
    "2.12.15": "Payment Policy", "2.12.16": "Forms of Violation, Disciplinary Actions", "2.12.17": "Disciplinary Actions",
    "2.12.18": "Procedural Due Process", "2.13": "Strategic Plan",
    "CHAPTER 3": "TECHNICAL ASPECT", "3.1": "Business Location", "3.1.1": "Office Location", "3.1.1.1": "Contract of Office Lease",
    "3.1.2": "Factory and Showroom Location", "3.1.2.1": "Contract of Factory and Showroom Lease", "3.1.3": "History of Location",
    "3.1.4": "Quick Facts", "3.1.5": "Barangays in Location", "3.1.6": "Geography", "3.1.7": "Economy", "3.2": "Floor Plan",
    "3.2.1": "Office Floor Plan", "3.2.2": "Factory and Showroom Floor Plan", "3.3": "List of Assets",
    "3.3.1": "List of Depreciable Assets", "3.3.2": "List of Non-Depreciable Assets", "3.4": "Product Description",
    "3.4.1": "Label or Packaging", "3.4.2": "Product", "3.4.2.1": "Table Cabinet", "3.4.2.2": "Expandable Table",
    "3.4.2.3": "Slot Sofa", "3.4.2.4": "Wall Bed", "3.5": "Production Process",
    "CHAPTER 4": "MARKETING ASPECT", "4.1": "Business Description", "4.2": "Industry Profile and Analysis",
    "4.3": "Study of Demand and Supply", "4.3.1": "Survey Participants Profile", "4.3.2": "Survey Results",
    "4.3.3": "Target Population", "4.3.4": "Projected Demand", "4.3.5": "Market Share", "4.3.6": "Projected Supply",
    "4.4": "Marketing Plan", "4.4.1": "Market Analysis", "4.4.1.1": "Market Demographics", "4.4.1.2": "Market Trend",
    "4.4.1.3": "Target Market", "4.4.1.4": "Competitors", "4.5": "Positioning", "4.6": "Marketing Mix", "4.6.1": "Product",
    "4.6.2": "Price", "4.6.3": "Advertising and Promotions", "4.6.3.1": "Online Advertising", "4.6.3.2": "Brochures and Catalogues",
    "4.6.4": "Placement",
    "CHAPTER 5": "SOCIAL AND ECONOMIC ASPECT", "5.1": "Economic Condition", "5.1.1": "Business in the Philippines",
    "5.1.2": "Business in the Proposed Location", "5.1.3": "The Specific Industry", "5.2": "Social Desirability",
    "5.2.1": "Personnel and Staff", "5.2.2": "Community", "5.2.3": "Skills Development", "5.2.4": "Suppliers",
    "5.2.5": "Investors", "5.2.6": "Political and Economic Contribution", "5.2.7": "Environmental Accountability",
    "5.2.8": "SWOT Analysis", "CHAPTER 6": "FINANCIAL ASPECT", "6.1": "Sources of Financing", "6.2": "Financial Assumptions",
    "6.3": "Capital Requirement", "6.4": "Projected Statement of Financial Position", "6.5": "Projected Statement of Comprehensive Income",
    "6.6": "Projected Statement of Changes in Equity", "6.7": "Projected Statement of Cash Flow", "6.8": "Notes to Financial Statements",
    "6.9": "Financial Ratios", "6.9.1": "Payback Period", "6.9.2": "Return on Investment", "6.9.3": "Horizontal Analysis",
    "6.9.4": "Vertical Analysis"
}

REVERSE_OUTLINE = {re.sub(r'[^a-zA-Z0-9]', '', v.lower()): k for k, v in FEASIBILITY_OUTLINE.items()}

# --- Deep XML Font Extractor ---
def get_deep_font_properties(p, doc):
    font_name, font_size = None, None
    for run in p.runs:
        if run.text.strip():
            if run.font.name: font_name = run.font.name
            if run.font.size: font_size = run.font.size.pt
            if not font_name or not font_size:
                try:
                    rPr = run._element.rPr
                    if rPr is not None:
                        if not font_name:
                            rFonts = rPr.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts')
                            if rFonts is not None:
                                theme_font = rFonts.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}asciiTheme')
                                ascii_font = rFonts.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ascii')
                                if theme_font: font_name = "Default Theme Font (Aptos/Calibri)"
                                elif ascii_font: font_name = ascii_font
                        if not font_size:
                            sz = rPr.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sz')
                            szCs = rPr.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}szCs')
                            if sz is not None:
                                val = sz.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')
                                if val: font_size = float(val) / 2
                            elif szCs is not None:
                                val = szCs.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')
                                if val: font_size = float(val) / 2
                except: pass
            if font_name and font_size: break
            
    if not font_name or not font_size:
        try:
            pPr = p._element.pPr
            if pPr is not None and pPr.rPr is not None:
                rPr = pPr.rPr
                if not font_name:
                    rFonts = rPr.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts')
                    if rFonts is not None:
                        theme_font = rFonts.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}asciiTheme')
                        if theme_font: font_name = "Default Theme Font (Aptos/Calibri)"
                if not font_size:
                    sz = rPr.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sz')
                    if sz is not None and sz.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val'):
                        font_size = float(sz.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')) / 2
        except: pass

    if not font_name or not font_size:
        curr_style = p.style
        while curr_style:
            if not font_name and curr_style.font.name: font_name = curr_style.font.name
            if not font_size and curr_style.font.size: font_size = curr_style.font.size.pt
            curr_style = curr_style.base_style

    if not font_name or not font_size:
        try:
            defaults = doc.styles.element.xpath('//w:docDefaults//w:rPr')
            if defaults:
                for element in defaults[0]:
                    if element.tag.endswith('rFonts') and not font_name:
                        if element.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}asciiTheme'): 
                            font_name = "Default Theme Font (Aptos/Calibri)"
                    if (element.tag.endswith('sz') or element.tag.endswith('szCs')) and not font_size:
                        val = element.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')
                        if val: font_size = float(val) / 2
        except: pass

    if not font_name: font_name = "Default Theme Font (Aptos/Calibri)"
    if not font_size: font_size = 11.0 

    return font_name, font_size

# --- Beautiful Web-Matched PDF Generator ---
class PDFReceipt(FPDF):
    def header(self):
        self.set_font('Times', 'B', 28)
        self.set_text_color(35, 55, 29) 
        self.cell(0, 10, 'Xeia Writes', ln=True, align='C')
        self.set_font('Times', 'I', 14)
        self.set_text_color(143, 179, 222) 
        self.cell(0, 6, 'a new academic paradigm', ln=True, align='C')
        self.set_draw_color(174, 167, 67) 
        self.line(10, 30, 200, 30)
        self.ln(10)

    def footer(self):
        self.set_y(-15)
        self.set_font('Helvetica', 'I', 8)
        self.set_text_color(128, 128, 128)
        self.cell(0, 10, f'Page {self.page_no()}', align='C')

    def add_centered_web_summary(self, f_val, s_val, total_paras):
        self.set_x(10)
        self.ln(5)
        y_start = self.get_y()
        
        cols = [
            {"x": 15, "title": "FORMATTING", "val": f"{f_val}%", "color": (143, 179, 222), "pct": f_val},
            {"x": 80, "title": "STRUCTURE", "val": f"{s_val}%", "color": (174, 167, 67), "pct": s_val},
            {"x": 145, "title": "SCOPE", "val": f"{total_paras}", "color": (69, 91, 48), "pct": 100}
        ]
        
        for col in cols:
            self.set_xy(col["x"], y_start)
            self.set_font('Helvetica', 'B', 8)
            self.set_text_color(69, 91, 48) 
            self.cell(50, 5, col["title"], ln=False)
            
            self.set_xy(col["x"], y_start + 6)
            self.set_font('Times', 'B', 24)
            self.set_text_color(35, 55, 29) 
            
            if col["title"] == "SCOPE":
                self.cell(10, 10, col["val"], ln=False)
                self.set_font('Times', '', 12)
                self.cell(30, 12, " paragraphs", ln=False)
            else:
                self.cell(50, 10, col["val"], ln=False)
            
            bar_y = y_start + 18
            bar_width = 50
            self.set_fill_color(230, 230, 230)
            self.rect(col["x"], bar_y, bar_width, 1.5, style='F')
            fill_width = (col["pct"] / 100) * bar_width
            r, g, b = col["color"]
            self.set_fill_color(r, g, b)
            self.rect(col["x"], bar_y, fill_width, 1.5, style='F')
            
        self.set_xy(10, y_start + 30)
        self.set_draw_color(220, 220, 220)
        self.line(10, self.get_y(), 200, self.get_y())
        self.ln(10)

def generate_pdf(active_lapses, f_score, s_score, total_paras):
    pdf = PDFReceipt()
    pdf.add_page()
    pdf.add_centered_web_summary(f_score, s_score, total_paras)
    
    section_titles = {
        "headings": "Heading & Outline Validations", "spelling": "Spelling & Typos", "grammar": "Basic Grammar", 
        "breaks": "Structural Breaks (Lines)", "ref_apa": "References: APA Format", "ref_indent": "References: Hanging Indents", 
        "ref_spacing": "References: Single Spacing", "font_name": "Font Style Deviations", "font_size": "Font Size Deviations", 
        "spacing": "Line Spacing Deviations", "indentation": "Indentation Deviations", "numbers": "Number Rule (0-9) Lapses",
        "suggestions": "Xeia-ggestions (Vocabulary & Flow)"
    }

    for cat_key, cat_title in section_titles.items():
        lapses_in_cat = active_lapses.get(cat_key, [])
        if lapses_in_cat:
            pdf.set_x(10)
            pdf.ln(5)
            
            pdf.set_font('Times', 'B', 18)
            if cat_key in ["breaks", "headings"]: pdf.set_text_color(143, 179, 222)
            elif cat_key == "suggestions": pdf.set_text_color(174, 167, 67)
            else: pdf.set_text_color(69, 91, 48)
            
            pdf.cell(0, 10, f'{cat_title}', ln=True)
            
            pdf.set_draw_color(220, 220, 220)
            pdf.line(10, pdf.get_y(), 200, pdf.get_y())
            pdf.ln(3)
            
            for lapse in lapses_in_cat:
                para, snippet, msg, _ = lapse
                
                clean_msg = msg.replace('<br>', '\n')
                clean_msg = re.sub(r'<[^>]+>', '', clean_msg).replace('➔', '->').replace('Applicability Check:', '\nApplicability Check:')
                clean_msg = clean_for_pdf(clean_msg)
                
                clean_snippet = snippet.replace('<br>', '\n')
                clean_snippet = re.sub(r'<[^>]+>', '', clean_snippet)
                clean_snippet = clean_for_pdf(clean_snippet)
                
                pdf.set_x(10)
                pdf.set_font('Helvetica', 'B', 11)
                pdf.set_text_color(69, 91, 48)
                pdf.cell(0, 6, f'Paragraph {para}:', ln=True)
                
                pdf.set_x(10)
                pdf.set_font('Helvetica', '', 10)
                pdf.set_text_color(0, 0, 0)
                pdf.multi_cell(0, 5, f'{clean_msg}')
                
                pdf.set_x(10)
                pdf.set_font('Helvetica', 'I', 9)
                pdf.set_text_color(143, 179, 222) 
                pdf.multi_cell(0, 5, f'"{clean_snippet}"')
                
                pdf.set_x(10)
                pdf.set_draw_color(230, 230, 230)
                pdf.line(10, pdf.get_y()+2, 200, pdf.get_y()+2)
                pdf.ln(5)

    if not any(active_lapses.values()):
        pdf.set_x(10)
        pdf.ln(10)
        pdf.set_font('Helvetica', 'B', 12)
        pdf.set_text_color(69, 91, 48)
        pdf.cell(0, 10, 'Flawless Execution! No active lapses found.', align='C')

    # THE FIX: Bulletproof encoding output to prevent Streamlit Cloud TypeError
    pdf_out = pdf.output(dest='S')
    return pdf_out.encode('latin-1') if isinstance(pdf_out, str) else bytes(pdf_out)

# --- Analysis Logic ---
def analyze_document(file, exp_font, exp_size, exp_spacing, exp_indent, number_rule, check_spelling, check_duplicates):
    doc = docx.Document(file)
    paragraphs = doc.paragraphs
    total_paras = 0
    in_references_section = False
    
    lapses = {
        "headings": [], "font_name": [], "font_size": [], "spacing": [], "indentation": [], 
        "breaks": [], "numbers": [], "spelling": [], "grammar": [], 
        "ref_spacing": [], "ref_indent": [], "ref_apa": [], "suggestions": []
    }
    
    spacing_map = {"Single": 1.0, "Double": 2.0, "1.5 lines": 1.5}
    target_spacing = spacing_map[exp_spacing]
    lapse_counter = [0]

    for i, p in enumerate(paragraphs):
        raw_text = p.text
        text = raw_text.strip()
            
        total_paras += 1
        para_num = i + 1 
        
        def add_lapse(category, msg, match_start=None, match_end=None, highlight=None):
            lapse_counter[0] += 1
            lapse_id = f"{category}_{para_num}_{lapse_counter[0]}_{uuid.uuid4().hex[:6]}"
            snippet = get_smart_snippet(text, match_start, match_end, highlight)
            lapses[category].append((para_num, snippet, msg, lapse_id))

        cleaned_text = text.lower()
        if cleaned_text in ["references", "bibliography", "works cited"] or (cleaned_text.endswith("references") and len(cleaned_text) < 20):
            in_references_section = True

        is_heading = False
        text_for_numbers = text
        heading_match = re.match(r'^(CHAPTER\s*\d+|[1-9](?:\.\d+[a-zA-Z]?)*)\s*[:\-]?\s*(.+)', text, re.IGNORECASE)
        
        if heading_match and not in_references_section:
            is_heading = True
            section_num = heading_match.group(1).upper().strip()
            section_num = re.sub(r'\s+', ' ', section_num) 
            section_title = heading_match.group(2).strip()
            clean_actual_title = re.sub(r'[^a-zA-Z0-9]', '', section_title.lower())
            
            text_for_numbers = section_title 
            
            if section_num in FEASIBILITY_OUTLINE:
                expected_title = FEASIBILITY_OUTLINE[section_num]
                clean_expected = re.sub(r'[^a-zA-Z0-9]', '', expected_title.lower())
                
                if clean_actual_title != clean_expected:
                    if clean_actual_title in REVERSE_OUTLINE:
                        correct_num = REVERSE_OUTLINE[clean_actual_title]
                        add_lapse("headings", f"Outline Mismatch: The section '{section_title}' should be numbered {correct_num}, not {section_num}.", heading_match.start(), heading_match.end(), highlight=text)
                    else:
                        add_lapse("headings", f"Outline Mismatch: Expected title '{expected_title}' for section {section_num}.", heading_match.start(), heading_match.end(), highlight=section_title)
            else:
                if clean_actual_title in REVERSE_OUTLINE:
                    correct_num = REVERSE_OUTLINE[clean_actual_title]
                    add_lapse("headings", f"Outline Mismatch: The section '{section_title}' should be numbered {correct_num}, not {section_num}.", heading_match.start(), heading_match.end(), highlight=text)
                else:
                    add_lapse("headings", f"Outline Mismatch: The section number '{section_num}' does not exist in the official syllabus.", heading_match.start(), heading_match.end(), highlight=section_num)
            
            is_bold = False
            is_italic = False
            
            for run in p.runs:
                if run.text.strip():
                    if run.bold: is_bold = True
                    if run.italic: is_italic = True
                    
            if not is_bold and p.style and p.style.font and p.style.font.bold:
                is_bold = True
            if not is_italic and p.style and p.style.font and p.style.font.italic:
                is_italic = True
                
            if not is_bold:
                add_lapse("headings", f"Formatting Error: The heading '{section_num} {section_title}' must be in Bold.", heading_match.start(), heading_match.end(), highlight=text)
            if is_italic:
                add_lapse("headings", f"Formatting Error: The heading '{section_num} {section_title}' contains italicized text, which is strictly prohibited.", heading_match.start(), heading_match.end(), highlight=text)

        et_al_trap = re.search(r'\bet\.\s*al\.', cleaned_text)
        if et_al_trap:
            add_lapse("ref_apa", "In-text Citation Error: Use 'et al.' (no period after 'et') instead of 'et. al.'", et_al_trap.start(), et_al_trap.end(), highlight="et. al.")
        
        missing_paren = re.search(r'\b([A-Z][a-z]+(?:\s+et\s+al\.?)?)\s+(\d{4})\b', text)
        if missing_paren:
            add_lapse("ref_apa", f"In-text Citation Error: APA requires parentheses around the year, e.g., '{missing_paren.group(1)} ({missing_paren.group(2)})'.", missing_paren.start(), missing_paren.end(), highlight=missing_paren.group())
            
        period_trap = re.search(r'[a-zA-Z0-9]\.\.(?!\.)', text)
        if period_trap:
            add_lapse("grammar", "Punctuation Error: Found double periods (..) where a single period is expected.", period_trap.start(), period_trap.end(), highlight="..")

        if not is_heading and not in_references_section:
            for pattern, (suggestion, context, critique) in xeia_lexicon.items():
                for match in re.finditer(pattern, text, re.IGNORECASE):
                    found_word = match.group(1) if len(match.groups()) > 0 else match.group(0)
                    msg = f"Instead of '{found_word}', consider using <b>{suggestion}</b>.<br><span style='font-size: 0.8rem; color:#455B30; font-style:italic;'><b>Xeia's Critique:</b> {critique}</span>"
                    add_lapse("suggestions", msg, match.start(), match.end(), highlight=found_word)

        if is_heading and not in_references_section:
            if '\n' in raw_text or '\x0b' in raw_text:
                 add_lapse("breaks", "Spacing Error: Found a manual soft break (Shift+Enter) inside title. Please remove it.")

        if in_references_section and not is_heading:
            if p.paragraph_format.line_spacing not in [1.0, 1]: add_lapse("ref_spacing", "Reference entry must be strictly Single Spaced.")
            if p.paragraph_format.first_line_indent is None or p.paragraph_format.first_line_indent.inches > -0.05: add_lapse("ref_indent", "Missing Hanging Indent (Highlight text > Right Click > Paragraph > Special: Hanging).")
            if not text.startswith("http"):
                if not re.search(r'\(\d{4}[a-z]?\)|\(n\.d\.?\)', text, re.IGNORECASE):
                    add_lapse("ref_apa", "APA format requires a year, e.g., (2024) or (n.d.).")
            continue 

        if number_rule:
            for match in re.finditer(r'\b[0-9]\b', text_for_numbers):
                offset = len(text) - len(text_for_numbers)
                add_lapse("numbers", f"Found single digit '{match.group()}'. Spell out numbers 0-9.", match.start() + offset, match.end() + offset, highlight=match.group())

        if check_spelling:
            words = re.findall(r'\b[a-zA-Z]+\b', text)
            unknown = spell.unknown(words)
            if unknown:
                for w in unknown:
                    suggestion = spell.correction(w)
                    if suggestion and suggestion != w: 
                        match = re.search(rf'\b{re.escape(w)}\b', text, re.IGNORECASE)
                        if match: add_lapse("spelling", f"Possible typo: <b>{w}</b> ➔ {suggestion}", match.start(), match.end(), highlight=w)

        if check_duplicates:
            for match in re.finditer(r'\b(\w+)\s+\1\b', text, re.IGNORECASE):
                add_lapse("grammar", f"Repeated word: '{match.group()}'", match.start(), match.end(), highlight=match.group())

        font_name_found, font_size_found = get_deep_font_properties(p, doc)

        if font_name_found != exp_font and not (font_name_found == "Times New Roman" and exp_font == "Times New Roman"):
            add_lapse("font_name", f"Found '{font_name_found}', expected '{exp_font}'")
        if font_size_found != exp_size: 
            add_lapse("font_size", f"Found size {font_size_found}, expected {exp_size}")

        spacing = p.paragraph_format.line_spacing
        if spacing is not None and spacing != target_spacing: add_lapse("spacing", f"Found {spacing} spacing, expected {target_spacing}")

        if exp_indent and not is_heading and len(text) > 50: 
            indent = p.paragraph_format.first_line_indent
            if indent is None or indent.inches == 0: add_lapse("indentation", "Missing first-line indent, expected 0.5 inch.")
            else:
                val_inches = round(indent.inches, 2)
                if abs(val_inches - 0.5) > 0.05: add_lapse("indentation", f"Found {val_inches} inch indent, expected 0.5 inch.")

    return total_paras, lapses

# --- Main App Logic ---
def main():
    st.markdown("""
    <style>
        @import url('https://fonts.googleapis.com/css2?family=Playfair+Display:ital,wght@0,500;0,700;1,500&family=Poppins:wght@300;400;500;600&display=swap');
        
        html { scroll-behavior: smooth; }
        html, body, [class*="css"] { font-family: 'Poppins', sans-serif; color: #455B30; background-color: #FAFCF7;}
        h1, h2, h3, h4 { font-family: 'Playfair Display', serif !important; color: #23371D !important; }
        
        [data-testid="stSidebar"] { display: none !important; }
        header { display: none !important; }
        .block-container { padding-top: 0rem !important; max-width: 1100px; }
        [data-testid="stStatusWidget"] { visibility: hidden !important; display: none !important; }
        
        .top-menu { display: flex; justify-content: space-between; align-items: center; padding: 15px 30px; font-size: 0.8rem; font-weight: 500; letter-spacing: 1px; color: #455B30; border-bottom: 1px solid #EAEAEA; background-color: rgba(250, 252, 247, 0.95); position: sticky; top: 0; z-index: 100; backdrop-filter: blur(5px);}
        .top-menu a { text-decoration: none; color: #455B30; transition: 0.2s; cursor: pointer; }
        .top-menu a:hover { color: #8FB3DE; }
        
        .hero-container { display: flex; flex-direction: column; align-items: center; text-align: center; width: 100%; margin-top: 30px;}
        .hero-logo-container img { height: 140px; mix-blend-mode: multiply; margin-bottom: 20px;}
        .hero-title { font-size: 4rem; font-weight: 700; margin-bottom: 0px; letter-spacing: -1px; color: #23371D;}
        .hero-subtitle { font-size: 1.1rem; color: #8FB3DE; font-family: 'Playfair Display', serif; font-style: italic; margin-top: 0; margin-bottom: 20px;}
        .hero-text { max-width: 600px; color: #455B30; line-height: 1.6; margin-bottom: 50px;}
        
        .action-box-container { border: 2px dashed #AEA743; border-radius: 20px; padding: 40px; background-color: rgba(250, 252, 247, 0.8); text-align: center; margin-bottom: 60px; position: relative;}
        .action-tag { font-family: 'Playfair Display', serif; background-color: #8FB3DE; color: white; padding: 5px 15px; border-radius: 20px; font-size: 0.75rem; font-weight: 600; position: absolute; top: -12px; left: 50%; transform: translateX(-50%); letter-spacing: 1px;}
        
        [data-testid="stExpander"] { background-color: white !important; border-radius: 10px !important; border: 1px solid #EAEAEA !important; box-shadow: 0 4px 10px rgba(0,0,0,0.02) !important; margin-bottom: 10px; text-align: left; }
        [data-testid="stExpander"] summary p { font-weight: 600 !important; color: #23371D !important; font-size: 0.95rem; }
        
        button[kind="primary"] { background-color: #AEA743 !important; color: white !important; border-radius: 30px; border: none; padding: 12px 35px; font-weight: 600; letter-spacing: 1px; transition: 0.3s; width: 100%; margin-top: 10px;}
        button[kind="primary"]:hover { background-color: #23371D !important; transform: translateY(-2px); box-shadow: 0 4px 12px rgba(35, 55, 29, 0.2); }
        
        button[kind="secondary"] { background-color: #FAFCF7 !important; color: #8FB3DE !important; border: 1px solid #8FB3DE !important; border-radius: 15px !important; padding: 2px 15px !important; font-size: 0.8rem !important; transition: 0.2s;}
        button[kind="secondary"]:hover { background-color: #8FB3DE !important; color: white !important; }
        
        .feature-card { background: white; border-radius: 15px; padding: 30px; text-align: center; box-shadow: 0 4px 15px rgba(0,0,0,0.03); height: 100%; border-top: 4px solid;}
        .feature-icon { font-size: 2.5rem; margin-bottom: 15px; }
        .feature-title { font-family: 'Playfair Display', serif; font-size: 1.3rem; font-weight: 700; color: #23371D; margin-bottom: 10px;}
        .feature-text { font-size: 0.85rem; color: #455B30; line-height: 1.5; }
        
        .finding-box { background-color: white; padding: 15px; margin-bottom: 10px; border-radius: 8px; border-left: 4px solid; box-shadow: 0 2px 5px rgba(0,0,0,0.02);}
        .mail-card { background: linear-gradient(145deg, #FFFDF5, #FAFCF7); padding: 15px; margin-bottom: 15px; border-radius: 10px; border: 1px solid #E2C785; box-shadow: 0 4px 10px rgba(174, 167, 67, 0.1); border-left: 5px solid #AEA743;}
        
        .metric-card { padding: 10px 5px; }
        .metric-title { font-size: 0.85rem; color: #455B30; font-weight: 600; text-transform: uppercase; margin-bottom: -5px;}
        .metric-value { font-family: 'Playfair Display', serif; font-size: 2.5rem; color: #23371D; line-height: 1.2;}
        .progress-bg { width: 100%; background-color: #EAEAEA; border-radius: 10px; height: 6px; margin-top: 5px; overflow: hidden;}
        
        .magic-dots { display: flex; justify-content: center; align-items: center; gap: 10px; margin-top: 15px; padding-bottom: 5px;}
        .magic-dots div { width: 12px; height: 12px; border-radius: 50%; background-color: #AEA743; animation: pixie-bounce 1.4s infinite ease-in-out both; box-shadow: 0 0 10px #AEA743; }
        .magic-dots div:nth-child(1) { animation-delay: -0.32s; background-color: #8FB3DE; box-shadow: 0 0 10px #8FB3DE;}
        .magic-dots div:nth-child(2) { animation-delay: -0.16s; background-color: #455B30; box-shadow: 0 0 10px #455B30;}
        @keyframes pixie-bounce { 0%, 80%, 100% { transform: scale(0.4); opacity: 0.3; } 40% { transform: scale(1.2); opacity: 1; } }
        .loading-label { text-align: center; font-family: 'Playfair Display', serif; font-style: italic; color: #23371D; margin-top: 5px; margin-bottom: 10px; font-size: 1.1rem; }
    </style>
    """, unsafe_allow_html=True)

    if 'ignored_lapses' not in st.session_state: st.session_state.ignored_lapses = set()
    if 'analysis_results' not in st.session_state: st.session_state.analysis_results = None
    if 'total_paras' not in st.session_state: st.session_state.total_paras = 0
    if 'open_lapses_category' not in st.session_state: st.session_state.open_lapses_category = None

    st.markdown("""
        <div class="top-menu">
            <a href="#about-section">(ABOUT)</a>
            <a href="#configuration-section">(CONFIGURATION)</a>
            <a href="#analysis-section">(ANALYSIS)</a>
            <a href="#features-section">(FEATURES)</a>
        </div>
    """, unsafe_allow_html=True)

    st.markdown("<div id='about-section'></div>", unsafe_allow_html=True)
    st.markdown("<div class='hero-container'>", unsafe_allow_html=True)
    try:
        st.markdown("<div class='hero-logo-container'>", unsafe_allow_html=True)
        st.image("logo.png", width=120)
        st.markdown("</div>", unsafe_allow_html=True)
    except:
        pass 
        
    st.markdown("""
        <h1 class="hero-title">Xeia Writes</h1>
        <p class="hero-subtitle">a new academic paradigm</p>
        <p class="hero-text">Join the journey towards better understanding your documents, beautifully formatting your feasibility studies, and unapologetically maintaining pristine APA references.</p>
        </div>
    """, unsafe_allow_html=True)

    st.markdown("<div id='configuration-section' style='padding-top: 20px;'></div>", unsafe_allow_html=True)
    st.markdown("""
        <div class="action-box-container">
            <div class="action-tag">BEFORE YOU SUBMIT</div>
            <h2 style='font-size: 1.8rem; margin-top:10px;'>Want to find out if your paper meets strict standards?!</h2>
            <p style='margin-bottom: 30px; font-size: 0.9rem;'>Configure your rules below and upload your document to begin.</p>
    """, unsafe_allow_html=True)

    col_cfg1, col_cfg2 = st.columns(2)
    with col_cfg1:
        with st.expander("✒️ Formatting Rules", expanded=False):
            expected_font = st.selectbox("Font Style", ["Times New Roman", "Arial", "Calibri"], key="cfg_font")
            expected_size = st.number_input("Font Size", min_value=8, max_value=24, value=12, key="cfg_size")
            spacing_rule = st.radio("Main Document Spacing", ["Single", "Double", "1.5 lines"], key="cfg_spacing")
            check_indent = st.checkbox("Standard Indents (0.5 inch)", value=True, key="cfg_indent")
    with col_cfg2:
        with st.expander("📜 Analysis Engines", expanded=False):
            number_rule = st.checkbox("0-9 in words, 10+ in digits", value=True, key="cfg_num")
            check_spelling = st.checkbox("Verify Spelling & Typos", value=True, key="cfg_spell")
            check_duplicates = st.checkbox("Flag Repeated Words", value=True, key="cfg_dup")
            st.info("APA Reference & Syllabus Outline checks activate automatically.")

    uploaded_file = st.file_uploader("", type=["docx"], key="file_upload")
    analyze_pressed = st.button("TAKE THE ANALYSIS", key="btn_analyze", type="primary")
    st.markdown("</div>", unsafe_allow_html=True) 

    st.markdown("<div id='analysis-section' style='padding-top: 20px;'></div>", unsafe_allow_html=True)
    
    if analyze_pressed:
        if uploaded_file is not None:
            loader_html = """
            <div style="padding: 15px; text-align: center; margin-bottom: 20px;">
                <div class="magic-dots"><div></div><div></div><div></div></div>
                <div class="loading-label">Synthesizing document...</div>
            </div>
            """
            loading_placeholder = st.empty()
            loading_placeholder.markdown(loader_html, unsafe_allow_html=True)
            
            paras, results = analyze_document(uploaded_file, expected_font, expected_size, spacing_rule, check_indent, number_rule, check_spelling, check_duplicates)
            st.session_state.total_paras = paras
            st.session_state.analysis_results = results
            st.session_state.ignored_lapses = set()
            st.session_state.open_lapses_category = None 
            loading_placeholder.empty()
        else:
            st.warning("Please upload a file first.")

    if st.session_state.analysis_results is not None:
        active_lapses = { "headings": [], "font_name": [], "font_size": [], "spacing": [], "indentation": [], "breaks": [], "numbers": [], "spelling": [], "grammar": [], "ref_spacing": [], "ref_indent": [], "ref_apa": [], "suggestions": [] }
        active_fmt_count = 0
        active_str_count = 0
        
        for cat, lapses in st.session_state.analysis_results.items():
            for lapse in lapses:
                if lapse[3] not in st.session_state.ignored_lapses:
                    active_lapses[cat].append(lapse)
                    if cat in ["headings", "font_name", "font_size", "spacing", "indentation", "breaks", "ref_spacing", "ref_indent", "ref_apa"]: active_fmt_count += 1
                    elif cat in ["numbers", "spelling", "grammar"]: active_str_count += 1

        total_fmt_checks = max(1, st.session_state.total_paras * 3) 
        total_str_checks = max(1, st.session_state.total_paras * 1.5)
        
        f_val = max(0, 100 - round((active_fmt_count / total_fmt_checks) * 100)) if st.session_state.total_paras > 0 else 0
        s_val = max(0, 100 - round((active_str_count / total_str_checks) * 100)) if st.session_state.total_paras > 0 else 0

        pdf_bytes = generate_pdf(active_lapses, f_val, s_val, st.session_state.total_paras)

        c_res1, c_res2, c_res3 = st.columns(3)
        with c_res1:
            st.markdown(f"""
            <div class="metric-card">
                <div class="metric-title">Formatting Conformance</div>
                <div class="metric-value">{f_val}%</div>
                <div class="progress-bg"><div style="width: {f_val}%; background-color: #8FB3DE; height: 100%;"></div></div>
            </div>
            """, unsafe_allow_html=True)
        with c_res2:
            st.markdown(f"""
            <div class="metric-card">
                <div class="metric-title">Language & Structure</div>
                <div class="metric-value">{s_val}%</div>
                <div class="progress-bg"><div style="width: {s_val}%; background-color: #AEA743; height: 100%;"></div></div>
            </div>
            """, unsafe_allow_html=True)
        with c_res3:
            st.markdown(f"""
            <div class="metric-card">
                <div class="metric-title">Document Scope</div>
                <div class="metric-value">{st.session_state.total_paras} <span style="font-size:1rem;">paragraphs</span></div>
                <div class="progress-bg"><div style="width: 100%; background-color: #455B30; height: 100%;"></div></div>
            </div>
            """, unsafe_allow_html=True)
            
        st.markdown("<hr style='border: 1px solid #EAEAEA;'>", unsafe_allow_html=True)

        col_lapses, col_side = st.columns([2.5, 1])

        with col_lapses:
            st.markdown("<h2 style='color:#23371D;'>Detailed Lapses Dashboard</h2>", unsafe_allow_html=True)
            
            def render_interactive_section(title, color, data, cat_key):
                is_open = (st.session_state.open_lapses_category == cat_key)
                
                with st.expander(f"{title} ({len(data)} items)", expanded=is_open):
                    for lapse in data:
                        para, snippet, msg, lapse_id = lapse
                        c_text, c_btn = st.columns([6, 1])
                        with c_text:
                            st.markdown(f"<div class='finding-box' style='border-left-color:{color};'><strong>Paragraph {para}:</strong> <span style='color: #23371D;'>{msg}</span><br><span style='color: #8FB3DE; font-size: 0.85rem;'>\"{snippet}\"</span></div>", unsafe_allow_html=True)
                        with c_btn:
                            st.markdown("<div style='height: 10px;'></div>", unsafe_allow_html=True) 
                            st.button("Ignore", key=f"btn_ign_{lapse_id}", on_click=ignore_lapse, args=(lapse_id, cat_key), use_container_width=True, type="secondary")

            if active_lapses["headings"]: render_interactive_section("Heading & Outline Validations", "#8CA364", active_lapses["headings"], "headings")
            if active_lapses["spelling"]: render_interactive_section("Spelling & Typos", "#C5C9BC", active_lapses["spelling"], "spelling")
            if active_lapses["grammar"]: render_interactive_section("Basic Grammar (Repetitions & Punctuation)", "#C5C9BC", active_lapses["grammar"], "grammar")
            if active_lapses["breaks"]: render_interactive_section("Structural Breaks (Lines)", "#8FB3DE", active_lapses["breaks"], "breaks")
            if active_lapses["ref_apa"]: render_interactive_section("References: APA Format", "#23371D", active_lapses["ref_apa"], "ref_apa")
            if active_lapses["ref_indent"]: render_interactive_section("References: Hanging Indents", "#23371D", active_lapses["ref_indent"], "ref_indent")
            if active_lapses["ref_spacing"]: render_interactive_section("References: Single Spacing", "#23371D", active_lapses["ref_spacing"], "ref_spacing")
            if active_lapses["font_name"]: render_interactive_section("Font Style Deviations", "#455B30", active_lapses["font_name"], "font_name")
            if active_lapses["font_size"]: render_interactive_section("Font Size Deviations", "#455B30", active_lapses["font_size"], "font_size")
            if active_lapses["spacing"]: render_interactive_section("Line Spacing Deviations", "#AEA743", active_lapses["spacing"], "spacing")
            if active_lapses["indentation"]: render_interactive_section("Indentation Deviations", "#AEA743", active_lapses["indentation"], "indentation")
            if active_lapses["numbers"]: render_interactive_section("Number Rule (0-9) Lapses", "#23371D", active_lapses["numbers"], "numbers")

            if (active_fmt_count + active_str_count) == 0 and st.session_state.total_paras > 0:
                 st.success("✨ Flawless execution! All rules are conformed to (or safely ignored).")

        with col_side:
            st.download_button(label="📥 Export PDF Receipt", data=pdf_bytes, file_name="Xeia_Writes_Receipt.pdf", mime="application/pdf", use_container_width=True, type="primary")
            
            if active_lapses["suggestions"]:
                st.markdown("<div style='margin-top: 15px;'></div>", unsafe_allow_html=True)
                
                sug_open = (st.session_state.open_lapses_category == "suggestions")
                with st.expander(f"💌 Xeia-ggestions for you ({len(active_lapses['suggestions'])})", expanded=sug_open):
                    for lapse in active_lapses["suggestions"]:
                        para, snippet, msg, lapse_id = lapse
                        st.markdown(f"""
                        <div class='mail-card'>
                            <strong>Paragraph {para}</strong><br>
                            <span style='color: #23371D; font-size: 0.9rem;'>{msg}</span><br>
                            <span style='color: #A0AAB2; font-size: 0.8rem; font-style:italic;'>"{snippet}"</span>
                        </div>
                        """, unsafe_allow_html=True)
                        st.button("Dismiss", key=f"btn_ign_{lapse_id}", on_click=ignore_lapse, args=(lapse_id, "suggestions"), use_container_width=True, type="secondary")

    else:
        st.markdown("<div id='features-section' style='padding-top: 40px;'></div>", unsafe_allow_html=True) 
        st.markdown("<h3 style='text-align: center; margin-bottom: 30px; font-size: 1.1rem; letter-spacing: 1px;'>LET'S FIGURE THIS ACADEMIC WRITING OUT TOGETHER</h3>", unsafe_allow_html=True)
        col_f1, col_f2, col_f3 = st.columns(3)
        with col_f1:
            st.markdown("""
            <div class="feature-card" style="border-top-color: #8FB3DE; background-color: #F4F7FC;">
                <div class="feature-icon">📏</div>
                <div style="font-size: 0.75rem; font-weight: 600; color: #8FB3DE; margin-bottom: 5px;">THE FORMATTING DECK</div>
                <div class="feature-title">Consistency Engine</div>
                <div class="feature-text">This is one of my favorite resources. I will help you identify rogue font sizes, broken indentations, and sneaky empty line breaks so that you can find peace of mind amidst the chaos of academic formatting.</div>
            </div>
            """, unsafe_allow_html=True)
        with col_f2:
            st.markdown("""
            <div class="feature-card" style="border-top-color: #AEA743; background-color: #FDFCEF;">
                <div class="feature-icon">🪶</div>
                <div style="font-size: 0.75rem; font-weight: 600; color: #AEA743; margin-bottom: 5px;">THE LANGUAGE DECK</div>
                <div class="feature-title">Structure & Flow</div>
                <div class="feature-text">Introducing the most comprehensive course to help you release your fears of typos. Inside this deck we peel back the layers of your grammar, identifying repeated words and offering brilliant vocabulary upgrades.</div>
            </div>
            """, unsafe_allow_html=True)
        with col_f3:
            st.markdown("""
            <div class="feature-card" style="border-top-color: #455B30; background-color: #F5F7F3;">
                <div class="feature-icon">📑</div>
                <div style="font-size: 0.75rem; font-weight: 600; color: #455B30; margin-bottom: 5px;">THE INTEGRITY DECK</div>
                <div class="feature-title">APA Validations</div>
                <div class="feature-text">Integrity is a transformational requirement for highly sensitive research. At its core, this deck will automatically detect your Bibliography and enforce strict hanging indents, single spacing, and publication years.</div>
            </div>
            """, unsafe_allow_html=True)

if __name__ == "__main__":
    main()
